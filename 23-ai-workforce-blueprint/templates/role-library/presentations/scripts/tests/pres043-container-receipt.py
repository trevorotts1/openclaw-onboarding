"""PRES-043 in-container acceptance receipt — runs INSIDE each Docker profile.

Renders the fixture deck with the profile's OWN interpreter and render chain
(PPTX -> LibreOffice PDF -> pdftoppm PNG -> tesseract OCR), then prints ONE
JSON receipt naming OS / interpreter / renderer / font versions and the
sha256 of every output. The Docker driver
(tests/pres043-matrix-docker.sh) captures this stdout as
release-matrix-receipt-<profile>.json, so the receipt can only ever describe
bytes this profile actually produced.

Usage (inside the profile image, workdir tests/):
    PROFILE=hostinger python3 pres043-container-receipt.py
"""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


def _sh(cmd: str) -> str:
    proc = subprocess.run(["/bin/bash", "-lc", cmd], capture_output=True,
                          text=True, timeout=120)
    return (proc.stdout + proc.stderr).strip()


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    import test_pres043_release_matrix as battery  # noqa: E402
    from pptx import Presentation  # noqa: E402
    from pptx.util import Inches  # noqa: E402

    profile = os.environ.get("PROFILE", "unknown")
    run_dir = Path("/tmp/pres043-in-container/probe")
    for sub in ("working/copy", "working/checkpoints", "working/deliverables",
                "working/qc", "working/units"):
        (run_dir / sub).mkdir(parents=True, exist_ok=True)

    pptx_path = run_dir / "PRES043-DECK01-FINAL.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(2)) \
            .text_frame.text = f"{battery.FIXTURE_COPY} slide {i}"
        slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12), Inches(2)) \
            .text_frame.text = (
                f"Fixture client {battery.FIXTURE_CLIENT} deck "
                f"{battery.DECK_IDS[0]} body line for OCR readback")
    prs.save(str(pptx_path))

    import pdf_export  # noqa: E402

    argv = sys.argv
    sys.argv = ["pdf_export.py", "--run-dir", str(run_dir)]
    try:
        pdf_export.main()
    finally:
        sys.argv = argv
    pdf_path = sorted((run_dir / "working" / "deliverables").glob("*.pdf"))[0]
    subprocess.run(
        ["pdftoppm", "-png", "-r", "150", "-f", "1", "-l", "1", str(pdf_path),
         str(run_dir / "working" / "deliverables" / "page")],
        check=True, capture_output=True, timeout=180)
    png_path = sorted((run_dir / "working" / "deliverables").glob("page-*.png"))[0]

    import pytesseract  # noqa: E402
    from PIL import Image  # noqa: E402

    ocr = pytesseract.image_to_string(Image.open(png_path)).strip()

    receipt = {
        "receipt_name": battery.RECEIPT_NAME,
        "profile": profile,
        "fixture_client": battery.FIXTURE_CLIENT,
        "deck_id": battery.DECK_IDS[0],
        "os": _sh("cat /etc/os-release | head -2 | tr '\\n' ';'"),
        "interpreter": _sh("python3 -c 'import sys; print(sys.version.split()[0], sys.executable)'"),
        "python": _sh("python3 --version"),
        "soffice": _sh("soffice --version | head -1"),
        "pdftoppm": _sh("pdftoppm -v 2>&1 | head -1"),
        "tesseract": _sh("tesseract --version 2>&1 | head -1"),
        "fonts_dejavu_liberation_faces": _sh("fc-list | grep -ciE 'dejavu|liberation'"),
        "fonts_total_faces": _sh("fc-list | wc -l"),
        "outputs": {
            "pptx": {"name": pptx_path.name, "sha256": _sha(pptx_path),
                     "bytes": pptx_path.stat().st_size},
            "pdf": {"name": pdf_path.name, "sha256": _sha(pdf_path),
                    "bytes": pdf_path.stat().st_size},
            "png": {"name": png_path.name, "sha256": _sha(png_path),
                    "bytes": png_path.stat().st_size},
        },
        "ocr_excerpt": ocr[:160],
    }
    print(json.dumps(receipt, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
