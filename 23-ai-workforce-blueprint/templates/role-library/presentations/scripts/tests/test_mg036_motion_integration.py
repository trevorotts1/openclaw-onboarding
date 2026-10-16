#!/usr/bin/env python3
"""test_mg036_motion_integration.py -- Skill 72 MG-036 integration tests.

MG-036 wires governed Skill 72 motion components (MG-035 packages) into
the existing presentation pipeline: slides.schema.json's typed media
object (P01), build_deck's motion authentication + embed (P02/P03/P04),
the P8.15-MOTION pipeline phase (P08), the webinar motion branch
(P05/P06), and deliverables.py's conditional motion components (P07).

The fixture is a REAL staged run: a hand-built MG-035-shaped package
whose component is a genuine silent 1920x1080 30 fps H.264 clip made
with ffmpeg, whose poster is the clip's final frame, and whose hashes
all bind the actual bytes on disk. Nothing here mocks the seam: the
validator, the planner, the assembler, the phase audit, the verifier,
and the webinar renderer all run against the staged bytes.

Conventions follow the house suite: pytest, SCRIPTS on sys.path,
tmp_path runs. Media-heavy tests skip when ffmpeg/ffprobe are absent.
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent
sys.path.insert(0, str(SCRIPTS))
REPO = SCRIPTS.parents[4]
SKILL72 = REPO / "72-motion-video-plus"
SCHEMA_FILES = [
    SCRIPTS / "slides.schema.json",
    REPO / "23-ai-workforce-blueprint" / "templates"
    / "presentation-render" / "slides.schema.json",
]
PACKAGE_SCHEMA = SKILL72 / "references" / "powerpoint-package-schema.json"
PIPELINE_MANIFEST = (REPO / "universal-sops" / "presentation-slide-craft"
                     / "PIPELINE-MANIFEST.json")

import pytest  # noqa: E402

import build_deck  # noqa: E402
import build_webinar_video as bwv  # noqa: E402
import motion_phase  # noqa: E402
import phase_verifiers  # noqa: E402
import webinar_ffmpeg as wf  # noqa: E402
from presentation_job import defers, deliverables, motion as pm  # noqa: E402

FFMPEG = shutil.which("ffmpeg")
FFPROBE = shutil.which("ffprobe")
needs_media = pytest.mark.skipif(
    not (FFMPEG and FFPROBE), reason="ffmpeg/ffprobe required")

JOB_ID = "job-mg036-fixture"
COPY_LINES = ["Northwind Co",
              "Three moves that doubled our pipeline in 90 days"]
DEFAULT_PLAYBACK = {"start": "on_click", "repeat": "once",
                    "end": "hold_final"}
LOOP_PLAYBACK = {"start": "automatic", "repeat": "loop", "end": "loop"}
FULL_PLACEMENT = {"x": 0.0, "y": 0.0, "width": 1.0, "height": 0.5625}

# The pre-MG-036 base the still-only regression compares against,
# pinned by explicit ref: on any tree where MG-036 is committed,
# git HEAD IS the edited pipeline, so an unpinned comparison would
# compare the pipeline against itself and pass vacuously. MG-048's
# branch update must re-point this ref if the branch is rebased
# past the MG-036 merge. The MOTION_REGRESSION_BASE environment
# variable wins when set.
REGRESSION_BASE_REF = "1d21d735b"


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sha256_file(p: Path) -> str:
    return _sha256_bytes(Path(p).read_bytes())


def _write_json(p: Path, obj) -> Path:
    p = Path(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, indent=2), encoding="utf-8")
    return p


def _run_ffmpeg(args):
    subprocess.run([FFMPEG, "-y", "-v", "error", *args],
                   check=True, capture_output=True, text=True)


def _probe_media(path: Path) -> dict:
    out = subprocess.run(
        [FFPROBE, "-v", "error", "-show_streams", "-show_format",
         "-of", "json", str(path)],
        check=True, capture_output=True, text=True).stdout
    info = json.loads(out)
    video = next(s for s in info["streams"] if s["codec_type"] == "video")
    num, den = video["avg_frame_rate"].split("/")
    return {
        "codec": video["codec_name"],
        "pixel_format": video["pix_fmt"],
        "width": video["width"],
        "height": video["height"],
        "frames": int(video.get("nb_frames") or 0),
        "fps": round(int(num) / int(den), 3),
        "duration_seconds": round(float(info["format"]["duration"]), 3),
        "audio_streams": sum(1 for s in info["streams"]
                             if s["codec_type"] == "audio"),
    }


def _make_component_clip(clip: Path) -> dict:
    """The canonical component: a 6.0 s silent 1920x1080 30 fps H.264
    clip on a bright ground: an ink bar sweeping left to right and an
    accent block drifting slower (overlay expressions, which ffmpeg
    evaluates per frame; drawbox geometry expressions are evaluated
    once and would leave the clip static). Returns the probe dict."""
    _run_ffmpeg([
        "-f", "lavfi", "-i", "color=c=0xF7F4EE:s=1920x1080:r=30:d=6",
        "-f", "lavfi", "-i", "color=c=0x192A38:s=420x160:r=30:d=6",
        "-f", "lavfi", "-i", "color=c=0xC04A26:s=260x80:r=30:d=6",
        "-filter_complex",
        ("[0:v][1:v]overlay=x='60+t*240':y=500[tmp];"
         "[tmp][2:v]overlay=x='60+t*90':y=200[out]"),
        "-map", "[out]",
        "-c:v", "libx264", "-preset", "veryfast", "-crf", "22",
        "-pix_fmt", "yuv420p", str(clip)])
    return _probe_media(clip)


def _overlay_filter() -> str:
    """The canonical component's filter graph (shared by the clip
    variants below so each variant differs from the canonical
    component in exactly the one property its test stages)."""
    return ("[0:v][1:v]overlay=x='60+t*240':y=500[tmp];"
            "[tmp][2:v]overlay=x='60+t*90':y=200[out]")


def _lavfi_layers(seconds: int) -> list:
    return [
        "-f", "lavfi", "-i",
        f"color=c=0xF7F4EE:s=1920x1080:r=30:d={seconds}",
        "-f", "lavfi", "-i",
        f"color=c=0x192A38:s=420x160:r=30:d={seconds}",
        "-f", "lavfi", "-i",
        f"color=c=0xC04A26:s=260x80:r=30:d={seconds}",
    ]


def _make_large_component_clip(clip: Path) -> dict:
    """The canonical component animation encoded as a 40 Mbps
    constant-bitrate broadcast export (HRD filler included): the
    governed H.264 shape in every respect (codec, pixel format,
    dimensions, rate, duration, silence), but its bytes run past
    the B14 20 MB delivery cap. Returns the probe dict."""
    _run_ffmpeg([
        *_lavfi_layers(6),
        "-filter_complex", _overlay_filter(),
        "-map", "[out]",
        "-c:v", "libx264", "-preset", "veryfast",
        "-b:v", "40M", "-minrate", "40M", "-maxrate", "40M",
        "-bufsize", "80M", "-x264-params", "nal-hrd=cbr",
        "-pix_fmt", "yuv420p", str(clip)])
    return _probe_media(clip)


def _make_short_component_clip(clip: Path) -> dict:
    """The canonical component at 3.0 seconds: governed shape,
    silent and small, but under the B14 4.0-second floor. Returns
    the probe dict."""
    _run_ffmpeg([
        *_lavfi_layers(3),
        "-filter_complex", _overlay_filter(),
        "-map", "[out]",
        "-c:v", "libx264", "-preset", "veryfast", "-crf", "22",
        "-pix_fmt", "yuv420p", str(clip)])
    return _probe_media(clip)


def _make_audio_component_clip(clip: Path) -> dict:
    """The canonical component with a sine tone muxed in as an AAC
    audio stream: governed shape and duration, but NOT silent.
    Returns the probe dict."""
    video_only = clip.with_name("video-only.mp4")
    _make_component_clip(video_only)
    _run_ffmpeg(["-i", str(video_only),
                 "-f", "lavfi", "-i", "sine=frequency=440:d=6",
                 "-map", "0:v", "-map", "1:a",
                 "-c:v", "copy", "-c:a", "aac", "-shortest",
                 str(clip)])
    return _probe_media(clip)


def _make_mpeg4_component_clip(clip: Path) -> dict:
    """The canonical component animation encoded with the MPEG-4
    Part 2 codec in an MP4 container: right dimensions, rate,
    duration and silence, wrong codec for the governed shape.
    Returns the probe dict."""
    _run_ffmpeg([
        *_lavfi_layers(6),
        "-filter_complex", _overlay_filter(),
        "-map", "[out]",
        "-c:v", "mpeg4", "-q:v", "4",
        "-pix_fmt", "yuv420p", str(clip)])
    return _probe_media(clip)


def _media_from_clip(clip: Path) -> dict:
    """Assemble the media dict build_run consumes from a finished
    clip: probe it and take its final frame as the poster (the same
    construction as the canonical media fixture)."""
    probe = _probe_media(clip)
    poster = clip.with_name("poster.png")
    _run_ffmpeg(["-i", str(clip),
                 "-vf", f"select='eq(n\\,{probe['frames'] - 1})'",
                 "-frames:v", "1", str(poster)])
    return {"clip": clip, "poster": poster, "probe": probe}


@pytest.fixture(scope="module")
def media(tmp_path_factory):
    """The component clip plus its final frame as the poster.
    Built once per module."""
    d = tmp_path_factory.mktemp("mg036-media")
    clip = d / "component.mp4"
    probe = _make_component_clip(clip)
    poster = d / "poster.png"
    _run_ffmpeg(["-i", str(clip),
                 "-vf", f"select='eq(n\\,{probe['frames'] - 1})'",
                 "-frames:v", "1", str(poster)])
    return {"clip": clip, "poster": poster, "probe": probe}


@pytest.fixture(scope="module")
def media_large(tmp_path_factory):
    """A component over the B14 20 MB byte cap, valid in every
    other respect. Built once per module."""
    d = tmp_path_factory.mktemp("mg036-media-large")
    clip = d / "component.mp4"
    _make_large_component_clip(clip)
    return _media_from_clip(clip)


@pytest.fixture(scope="module")
def media_short(tmp_path_factory):
    """A component under the B14 4.0-second duration floor, valid
    in every other respect. Built once per module."""
    d = tmp_path_factory.mktemp("mg036-media-short")
    clip = d / "component.mp4"
    _make_short_component_clip(clip)
    return _media_from_clip(clip)


@pytest.fixture(scope="module")
def media_audio(tmp_path_factory):
    """A component carrying an audio stream, valid in every other
    respect. Built once per module."""
    d = tmp_path_factory.mktemp("mg036-media-audio")
    clip = d / "component.mp4"
    _make_audio_component_clip(clip)
    return _media_from_clip(clip)


@pytest.fixture(scope="module")
def media_mpeg4(tmp_path_factory):
    """A component encoded with the wrong codec, valid in every
    other respect. Built once per module."""
    d = tmp_path_factory.mktemp("mg036-media-mpeg4")
    clip = d / "component.mp4"
    _make_mpeg4_component_clip(clip)
    return _media_from_clip(clip)


def media_block(slide_no: int, *, fit="hold_final",
                playback=None, placement=None) -> dict:
    """The P01 slides.json media object for one staged slide."""
    ws = f"working/motion/slide-{slide_no:02d}"
    return {
        "kind": "motion",
        "motion_manifest_path": f"{ws}/manifest.json",
        "approved_render_receipt_path": f"{ws}/renders/artifacts.json",
        "poster_path": f"{ws}/delivery/poster.png",
        "placement": dict(placement or FULL_PLACEMENT),
        "playback": dict(playback or DEFAULT_PLAYBACK),
        "webinar_policy": {"fit": fit, "duplicate_audio": False},
    }


def build_package(run_dir: Path, slide_no: int, media: dict, *,
                  qc_verdict="pass", checks_failed=0,
                  playback=None, hold_status="pass",
                  loop_seam_status="not_applicable",
                  essential_text=None, receipt_status="approved",
                  receipt_artifact_sha=None,
                  storyboard_strings=None,
                  poster_bytes=None) -> dict:
    """Stage one governed MG-035 package under the run and return its
    hashes. Every hash binds the real staged bytes; overrides exist so
    the degrade tests can break exactly one link at a time."""
    playback = dict(playback or DEFAULT_PLAYBACK)
    ws = Path(run_dir) / "working" / "motion" / f"slide-{slide_no:02d}"
    pkg = ws / "package"
    pkg.mkdir(parents=True, exist_ok=True)
    clip_bytes = media["clip"].read_bytes()
    (pkg / "component.mp4").write_bytes(clip_bytes)
    clip_sha = _sha256_bytes(clip_bytes)
    poster_data = (poster_bytes if poster_bytes is not None
                   else media["poster"].read_bytes())
    poster_sha = _sha256_bytes(poster_data)
    (pkg / "poster.png").write_bytes(poster_data)
    (pkg / "fallback.png").write_bytes(poster_data)
    (ws / "delivery").mkdir(parents=True, exist_ok=True)
    (ws / "delivery" / "poster.png").write_bytes(poster_data)

    manifest = {"job_id": JOB_ID, "mode": "powerpoint_graphic",
                "delivery": {"video_profile": "h264_opaque",
                             "playback": playback}}
    manifest_path = _write_json(ws / "manifest.json", manifest)

    source = pkg / "source"
    storyboard = {"beats": [{"text": s} for s in
                            (storyboard_strings or COPY_LINES)]}
    src_paths = {
        "manifest": _write_json(source / "manifest.json", manifest),
        "timeline": _write_json(source / "timeline.json",
                                {"frames": media["probe"]["frames"]}),
        "sidecar": _write_json(source / "sidecar.json",
                               {"job_id": JOB_ID}),
        "storyboard": _write_json(source / "storyboard.json", storyboard),
        "design": _write_json(source / "design.json",
                              {"background": "#F7F4EE"}),
    }
    qc_doc = {"job_id": JOB_ID, "verdict": qc_verdict,
              "artifact_sha256": clip_sha}
    qc_path = _write_json(pkg / "final-qc.json", qc_doc)

    probe = dict(media["probe"])
    record = {
        "schema_version": "1.0",
        "item": "MG-035",
        "job_id": JOB_ID,
        "mode": "powerpoint_graphic",
        "generated_by": "test_mg036_motion_integration fixture builder",
        "source_job": {
            "manifest_sha256": _sha256_file(manifest_path),
            "timeline_sha256": _sha256_file(src_paths["timeline"]),
            "sidecar_sha256": _sha256_file(src_paths["sidecar"]),
            "storyboard_sha256": _sha256_file(src_paths["storyboard"]),
            "design_sha256": _sha256_file(src_paths["design"]),
        },
        "artifact": {"file": "component.mp4", "sha256": clip_sha,
                     "bytes": len(clip_bytes), "probe": probe},
        "export": {"profile": "h264_opaque",
                   "export_record_sha256": None},
        "audio": {"resolved_mode": "none", "policy_version": "mg023-1",
                  "artifact_audio_streams": 0},
        "poster": {
            "file": "poster.png", "sha256": poster_sha,
            "bytes": len(poster_data), "width": 1920, "height": 1080,
            "frame": probe["frames"] - 1, "state": "final_hold",
            "rule": "final_hold_start",
            "reason": "the final frame is the readable hold state",
            "stable_interval": [probe["frames"] - 30, probe["frames"]],
            "motion_complete_frame": probe["frames"] - 30,
            "declared_poster_frame": None, "per_scene": [],
            "measured": None,
        },
        "fallback": {"file": "fallback.png", "sha256": poster_sha,
                     "bytes": len(poster_data),
                     "source_frame": probe["frames"] - 1,
                     "carries": "the poster state as a still"},
        "playback": {
            **playback,
            "intended_duration_frames": probe["frames"],
            "intended_duration_seconds": probe["duration_seconds"],
            "final_hold_state": "the readable final state",
        },
        "placement": {
            "kind": "full_bleed", "display_width_px": 1920,
            "display_height_px": 1080, "slide_width_in": 10.0,
            "slide_height_in": 5.625, "scale": 1.0, "roles": [],
            "essential_text": list(essential_text or COPY_LINES),
            "caveat": "fixture placement",
        },
        "checks": {
            "loop_seam": {"status": loop_seam_status},
            "hold_final": {"status": hold_status},
        },
        "qc": {"record_file": "final-qc.json",
               "record_sha256": _sha256_file(qc_path),
               "verdict": qc_verdict, "artifact_sha256": clip_sha,
               "checks_passed": 27 if qc_verdict == "pass" else 26,
               "checks_failed": checks_failed},
        "deck": {"file": "fixture-FINAL.pptx", "sha256": clip_sha,
                 "bytes": 1, "slides": 3,
                 "media_part": "ppt/media/media1.mp4",
                 "embed_verified": True},
        "platform": {"powerpoint_playback": "UNVERIFIED",
                     "statement": "no PowerPoint playback claim"},
    }
    _write_json(pkg / "package-record.json", record)

    receipt = {"job_id": JOB_ID,
               "artifact_sha256": receipt_artifact_sha or clip_sha,
               "status": receipt_status}
    _write_json(ws / "renders" / "artifacts.json", receipt)
    return {"workspace": ws, "package_dir": pkg, "record": record,
            "clip_sha": clip_sha, "poster_sha": poster_sha}


def _still_png(path: Path, seed: int) -> Path:
    """A photographic-style still render (ffmpeg test pattern): low
    flat-fill dominance, comfortably over the byte floor."""
    path.parent.mkdir(parents=True, exist_ok=True)
    _run_ffmpeg(["-f", "lavfi", "-i",
                 f"testsrc2=s=1280x720:d=1:r=1",
                 "-frames:v", "1", "-update", "1", str(path)])
    return path


def build_run(tmp_path: Path, media: dict, *, elected=True,
              motion_slides=(1,), package_overrides=None,
              arc_dividers=(), poster_bytes=None,
              extra_motion_slide_nos=()) -> Path:
    """Stage a 3-slide run directory: intake, slides.json, renders,
    and one governed package per motion slide."""
    run_dir = tmp_path / "run"
    (run_dir / "working" / "copy").mkdir(parents=True, exist_ok=True)
    (run_dir / "renders").mkdir(parents=True, exist_ok=True)
    intake = {}
    if elected:
        intake = {"pre_presentation_capture":
                  {"WANT_MOTION_GRAPHICS": "yes"}}
    _write_json(run_dir / "working" / "copy" / "intake.json", intake)
    if arc_dividers:
        arc = {"slides": [
            {"slide": n, "arc_section": "section divider",
             "tag": "divider"} for n in arc_dividers]}
        _write_json(run_dir / "working" / "copy" / "arc_allocation.json",
                    arc)
    overrides = package_overrides or {}
    slides = []
    for n in (1, 2, 3):
        slide = {"slide": n, "scene": f"scene-{n}",
                 "copy": list(COPY_LINES) if n == 1
                 else [f"Slide {n} headline", f"Slide {n} support"]}
        if n in motion_slides or n in extra_motion_slide_nos:
            slide["media"] = media_block(n, **overrides.get(
                "media", {}))
        slides.append(slide)
    _write_json(run_dir / "working" / "copy" / "slides.json", slides)
    for n in motion_slides:
        build_package(run_dir, n, media,
                      poster_bytes=poster_bytes,
                      **overrides.get("package", {}))
    for n in (1, 2, 3):
        png = run_dir / "renders" / f"slide-{n:02d}.png"
        if n in motion_slides:
            data = (poster_bytes if poster_bytes is not None
                    else media["poster"].read_bytes())
            png.write_bytes(data)
        else:
            _still_png(png, n)
    return run_dir


def _slides_of(run_dir: Path):
    return json.loads(
        (run_dir / "working" / "copy" / "slides.json").read_text())


def _outcome_for(record: dict, slide_no: int) -> dict:
    return next(e for e in record["slides"] if e["slide"] == slide_no)


# ---------------------------------------------------------------------------
# P01: the typed media object in slides.schema.json (both copies)
# ---------------------------------------------------------------------------

def _draft7():
    jsonschema = pytest.importorskip("jsonschema")
    return jsonschema.Draft7Validator


@pytest.mark.parametrize("schema_path", SCHEMA_FILES,
                         ids=["scripts", "presentation-render"])
def test_p01_schema_media_shapes(schema_path):
    validator_cls = _draft7()
    schema = json.loads(Path(schema_path).read_text())
    v = validator_cls(schema)
    base = {"slide": 1, "scene": "hook", "copy": ["Northwind Co"]}
    assert v.is_valid([base]), (
        "a slide with no media is the still default")
    assert v.is_valid([{**base, "media": {"kind": "still"}}])
    motion = {**base, "media": media_block(1)}
    assert v.is_valid([motion]), list(v.iter_errors([motion]))
    missing = media_block(1)
    del missing["poster_path"]
    assert not v.is_valid([{**base, "media": missing}]), (
        "a motion slide without its poster path must not validate")
    extra = media_block(1)
    extra["surprise"] = True
    assert not v.is_valid([{**base, "media": extra}]), (
        "the motion object is closed (additionalProperties false)")
    bad_fit = media_block(1)
    bad_fit["webinar_policy"] = {"fit": "stretch", "duplicate_audio": False}
    assert not v.is_valid([{**base, "media": bad_fit}])


def test_p01_both_schema_copies_carry_identical_media():
    blocks = []
    for p in SCHEMA_FILES:
        schema = json.loads(Path(p).read_text())
        blocks.append(schema["items"]["properties"]["media"])
    assert blocks[0] == blocks[1], (
        "the two slides.schema.json copies must stay in lockstep")


def test_fixture_package_record_satisfies_mg035_schema(media, tmp_path):
    validator_cls = _draft7()
    schema = json.loads(PACKAGE_SCHEMA.read_text())
    run_dir = build_run(tmp_path, media)
    record = json.loads((run_dir / "working" / "motion" / "slide-01"
                         / "package" / "package-record.json").read_text())
    errors = list(validator_cls(schema).iter_errors(record))
    assert errors == [], [e.message for e in errors]


# ---------------------------------------------------------------------------
# Election (B14): default NO, capture YES, waiver NO
# ---------------------------------------------------------------------------

def test_election_default_no_and_capture_yes(tmp_path):
    run_dir = tmp_path / "run"
    (run_dir / "working" / "copy").mkdir(parents=True)
    _write_json(run_dir / "working" / "copy" / "intake.json", {})
    e = pm.election_for_run(run_dir)
    assert e["elected"] is False and e["value"] == "no"
    assert e["source"] == "default"
    _write_json(run_dir / "working" / "copy" / "intake.json",
                {"pre_presentation_capture":
                 {"WANT_MOTION_GRAPHICS": "yes"}})
    e2 = pm.election_for_run(run_dir)
    assert e2["elected"] is True and e2["source"] == "intake_capture"
    # The registration lives in defers' want-key table with the same
    # semantics as the house elections.
    assert defers.resolve_intake_value(
        {"pre_presentation_capture": {}}, "want_motion_graphics") == "no"


# ---------------------------------------------------------------------------
# P02: resolve + validate
# ---------------------------------------------------------------------------

def test_resolve_slide_media_defaults_and_malformed(tmp_path):
    slides = [{"slide": 1},
              {"slide": 2, "media": {"kind": "still"}},
              {"slide": 3, "media": "not-an-object"}]
    out = build_deck.resolve_slide_media(slides, tmp_path)
    assert out[1] == {"kind": "still"}
    assert out[2] == {"kind": "still"}
    assert out[3]["kind"] == "motion" and out[3]["_malformed"] is True


@needs_media
def test_validate_happy_path(media, tmp_path):
    run_dir = build_run(tmp_path, media)
    slides = _slides_of(run_dir)
    election = pm.election_for_run(run_dir)
    outcome = build_deck.validate_motion_slide_artifacts(
        run_dir, 1, slides[0]["media"], election=election,
        copy_lines=list(COPY_LINES))
    assert outcome["status"] == "validated", outcome
    assert outcome["slot"] == "A"
    assert outcome["job_id"] == JOB_ID
    assert outcome["artifact_sha256"] == _sha256_file(media["clip"])
    assert outcome["poster_sha256"] == _sha256_file(media["poster"])
    assert outcome["_embed"]["clip_path"].name == "component.mp4"


@needs_media
def test_validate_degrades_by_name(media, tmp_path):
    """Every staged failure degrades to static with its own reason
    code: the acceptance's named degrade cases plus the rest of the
    vocabulary the seam promises."""
    cases = {}

    rd = build_run(tmp_path / "a", media, elected=False)
    cases["not_elected"] = (rd, pm.MOTION_NOT_ELECTED)

    rd = build_run(tmp_path / "b", media)
    shutil.rmtree(rd / "working" / "motion" / "slide-01" / "package")
    cases["missing_package"] = (rd, pm.MOTION_PACKAGE_MISSING)

    rd = build_run(tmp_path / "c", media, package_overrides={
        "package": {"qc_verdict": "fail", "checks_failed": 1}})
    cases["failed_qc"] = (rd, pm.MOTION_QC_FAILED)

    rd = build_run(tmp_path / "d", media)
    poster = rd / "working" / "motion" / "slide-01" / "delivery" \
        / "poster.png"
    poster.write_bytes(poster.read_bytes() + b"tampered")
    cases["poster_mismatch"] = (rd, pm.MOTION_POSTER_MISMATCH)

    rd = build_run(tmp_path / "e", media, package_overrides={
        "package": {"receipt_artifact_sha": "0" * 64}})
    cases["receipt_mismatch"] = (rd, pm.MOTION_RECEIPT_MISMATCH)

    rd = build_run(tmp_path / "f", media)
    slides_f = _slides_of(rd)
    cases["slot_invalid"] = (
        (rd, slides_f[0]["media"], 2), pm.MOTION_SLOT_INVALID)

    rd = build_run(tmp_path / "g", media)
    bad_playback = media_block(1, playback={
        "start": "automatic", "repeat": "once", "end": "hold_final"})
    cases["playback_mismatch"] = ((rd, bad_playback, 1),
                                  pm.MOTION_PLAYBACK_MISMATCH)

    rd = build_run(tmp_path / "h", media, package_overrides={
        "media": {"fit": "approved_loop"}})
    cases["loop_unqualified"] = (rd, pm.MOTION_LOOP_UNQUALIFIED)

    rd = build_run(tmp_path / "i", media, package_overrides={
        "package": {"storyboard_strings": ["Unrelated text"],
                    "essential_text": ["Unrelated text"]}})
    cases["copy_inventory"] = (rd, pm.MOTION_COPY_INVENTORY_INCOMPLETE)

    for name, payload in cases.items():
        if name in ("slot_invalid", "playback_mismatch"):
            (rd, media_obj, ord_), expected = payload
        else:
            rd, expected = payload
            media_obj, ord_ = _slides_of(rd)[0]["media"], 1
        election = pm.election_for_run(rd)
        outcome = build_deck.validate_motion_slide_artifacts(
            rd, ord_, media_obj, election=election,
            copy_lines=list(COPY_LINES) if ord_ == 1 else None)
        assert outcome["status"] == "degraded", (name, outcome)
        assert outcome["reason_code"] == expected, (
            name, outcome["reason_code"], outcome["detail"])
        assert "_embed" not in outcome, name


def _validate_only(media_variant, tmp_path) -> dict:
    """Stage a run whose single motion slide carries the variant
    component and validate it: the shared shape of the clip-limit
    tests below, each of which stages exactly one violation against
    an otherwise valid package."""
    run_dir = build_run(tmp_path, media_variant)
    slides = _slides_of(run_dir)
    return build_deck.validate_motion_slide_artifacts(
        run_dir, 1, slides[0]["media"],
        election=pm.election_for_run(run_dir),
        copy_lines=list(COPY_LINES))


@needs_media
def test_validate_clip_too_large_degrades(media_large, tmp_path):
    """One staged violation: the component's bytes exceed the B14
    20 MB cap. The package is valid in every other respect, so the
    validator must refuse with MOTION_CLIP_TOO_LARGE and the slide
    degrades (no embed)."""
    assert media_large["clip"].stat().st_size \
        > pm.MOTION_CLIP_MAX_BYTES, (
            "the staged fixture must actually exceed the cap it "
            "tests")
    outcome = _validate_only(media_large, tmp_path)
    assert outcome["status"] == "degraded", outcome
    assert outcome["reason_code"] == pm.MOTION_CLIP_TOO_LARGE, outcome
    assert "_embed" not in outcome


@needs_media
def test_validate_clip_duration_degrades(media_short, tmp_path):
    """One staged violation: the component runs 3.0 seconds, under
    the B14 4.0-second floor. Otherwise valid, so the refusal is
    MOTION_CLIP_DURATION and the slide degrades (no embed)."""
    assert media_short["probe"]["duration_seconds"] \
        < pm.MOTION_CLIP_MIN_SECONDS, (
            "the staged fixture must actually break the floor it "
            "tests")
    outcome = _validate_only(media_short, tmp_path)
    assert outcome["status"] == "degraded", outcome
    assert outcome["reason_code"] == pm.MOTION_CLIP_DURATION, outcome
    assert "_embed" not in outcome


@needs_media
def test_validate_clip_not_silent_degrades(media_audio, tmp_path):
    """One staged violation: the component carries an audio stream.
    A deck motion piece is silent (the presenter narrates), so the
    refusal is MOTION_CLIP_NOT_SILENT and the slide degrades (no
    embed)."""
    assert media_audio["probe"]["audio_streams"] == 1, (
        "the staged fixture must actually carry an audio stream")
    run_dir = build_run(tmp_path, media_audio)
    # The package record must describe its own artifact truthfully:
    # build_package's default record declares the silent fixture
    # shape, so restate its audio block for this staged package.
    record_path = (run_dir / "working" / "motion" / "slide-01"
                   / "package" / "package-record.json")
    record = json.loads(record_path.read_text())
    record["audio"]["artifact_audio_streams"] = 1
    _write_json(record_path, record)
    slides = _slides_of(run_dir)
    outcome = build_deck.validate_motion_slide_artifacts(
        run_dir, 1, slides[0]["media"],
        election=pm.election_for_run(run_dir),
        copy_lines=list(COPY_LINES))
    assert outcome["status"] == "degraded", outcome
    assert outcome["reason_code"] == pm.MOTION_CLIP_NOT_SILENT, outcome
    assert "_embed" not in outcome


@needs_media
def test_validate_clip_format_degrades(media_mpeg4, tmp_path):
    """One staged violation: the component is MPEG-4 Part 2, not
    the governed H.264 shape. Dimensions, rate, duration and
    silence all conform, so the refusal is MOTION_CLIP_FORMAT and
    the slide degrades (no embed)."""
    assert media_mpeg4["probe"]["codec"] != "h264", (
        "the staged fixture must actually carry the wrong codec")
    assert media_mpeg4["probe"]["width"] == pm.MOTION_CLIP_WIDTH
    assert media_mpeg4["probe"]["height"] == pm.MOTION_CLIP_HEIGHT
    outcome = _validate_only(media_mpeg4, tmp_path)
    assert outcome["status"] == "degraded", outcome
    assert outcome["reason_code"] == pm.MOTION_CLIP_FORMAT, outcome
    assert "_embed" not in outcome


@needs_media
def test_validate_slot_b_section_divider(media, tmp_path):
    run_dir = build_run(tmp_path, media, motion_slides=(3,),
                        arc_dividers=(3,))
    slides = _slides_of(run_dir)
    election = pm.election_for_run(run_dir)
    outcome = build_deck.validate_motion_slide_artifacts(
        run_dir, 3, slides[2]["media"], election=election,
        arc_tag_blob="section divider", copy_lines=[])
    assert outcome["status"] == "validated", outcome
    assert outcome["slot"] == "B"


# ---------------------------------------------------------------------------
# Plan: records, cap, still-only behavior
# ---------------------------------------------------------------------------

@needs_media
def test_plan_writes_records_and_manifest(media, tmp_path):
    run_dir = build_run(tmp_path, media)
    record, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    assert set(motion_media) == {1}
    entry = _outcome_for(record, 1)
    assert entry["status"] == "validated" and entry["slot"] == "A"
    on_disk = json.loads((run_dir / pm.MOTION_RECORD_REL).read_text())
    assert on_disk["election"]["elected"] is True
    manifest = json.loads(
        (run_dir / pm.MOTION_SLIDE_MEDIA_MANIFEST_REL).read_text())
    assert [e["slide"] for e in manifest["slides"]] == [1]
    assert manifest["slides"][0]["clip_sha256"] == entry["artifact_sha256"]
    assert manifest["slides"][0]["fit"] == "hold_final"


@needs_media
def test_plan_enforces_b14_cap(media, tmp_path):
    all_copy = list(COPY_LINES) + [
        "Slide 2 headline", "Slide 2 support",
        "Slide 3 headline", "Slide 3 support"]
    run_dir = build_run(
        tmp_path, media, motion_slides=(1, 2, 3), arc_dividers=(2, 3),
        package_overrides={"package": {
            "storyboard_strings": all_copy,
            "essential_text": list(COPY_LINES)}})
    record, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    assert set(motion_media) == {1, 2}, (
        "the B14 hard cap is 2 pieces per deck")
    third = _outcome_for(record, 3)
    assert third["status"] == "degraded"
    assert third["reason_code"] == pm.MOTION_CAP_EXCEEDED


@needs_media
def test_plan_still_only_deck_is_quiet(media, tmp_path):
    run_dir = build_run(tmp_path, media, motion_slides=())
    record, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    assert motion_media == {}
    assert record["slides"] == []
    assert record["election"]["elected"] is True
    manifest = json.loads(
        (run_dir / pm.MOTION_SLIDE_MEDIA_MANIFEST_REL).read_text())
    assert manifest["slides"] == []


# ---------------------------------------------------------------------------
# P04: assembly embeds through add_motion_media; notes preserved;
# re-assembly never duplicates media
# ---------------------------------------------------------------------------

def _deck_parts(deck: Path) -> dict:
    with zipfile.ZipFile(deck) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _timing_signature(slide_xml: bytes):
    import xml.etree.ElementTree as ET
    root = ET.fromstring(slide_xml)
    keep = ("delay", "repeatCount", "fill", "presetClass", "presetID",
            "presetSubtype", "dur", "nodeType")
    sig = []
    for el in root.iter():
        name = _local(el.tag)
        if name in ("timing", "tnLst", "par", "cTn", "stCondLst", "cond",
                    "cMediaNode", "tgtEl", "spTgt", "childTnLst", "seq",
                    "prevCondLst", "nextCondLst"):
            attrs = tuple(sorted(
                (k, v) for k, v in el.attrib.items() if k in keep))
            sig.append((name, attrs))
    return sig


def _media_node_ctn(slide_xml: bytes):
    import xml.etree.ElementTree as ET
    root = ET.fromstring(slide_xml)
    for el in root.iter():
        if _local(el.tag) == "cMediaNode":
            for child in el:
                if _local(child.tag) == "cTn":
                    return child
    return None


def _rendered_list(run_dir: Path):
    return [{"slide": n,
             "file": str(run_dir / "renders" / f"slide-{n:02d}.png")}
            for n in (1, 2, 3)]


SPEECH = {1: "Notes for slide one.", 2: "Notes for slide two.",
          3: "Notes for slide three."}


@needs_media
def test_assemble_embeds_movie_and_is_idempotent(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(tmp_path, media)
    _record, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    deck = tmp_path / "deck.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), deck,
                             speech_chunks=SPEECH,
                             motion_media=motion_media)
    parts = _deck_parts(deck)
    media_parts = [n for n in parts if n.startswith("ppt/media/")
                   and n.lower().endswith((".mp4", ".mov", ".m4v",
                                           ".webm"))]
    assert len(media_parts) == 1, media_parts
    assert parts[media_parts[0]] == media["clip"].read_bytes(), (
        "the embedded media part must be byte-identical to the "
        "authenticated component")

    import xml.etree.ElementTree as ET
    slide1 = parts["ppt/slides/slide1.xml"]
    tags = {_local(el.tag) for el in ET.fromstring(slide1).iter()}
    assert "videoFile" in tags and "media" in tags, (
        "slide 1 must carry a:videoFile and the p14 media element")
    tags2 = {_local(el.tag)
             for el in ET.fromstring(parts["ppt/slides/slide2.xml"]).iter()}
    assert "videoFile" not in tags2, "a still slide gains no movie"

    rels = parts["ppt/slides/_rels/slide1.xml.rels"].decode()
    assert 'TargetMode="External"' not in rels, (
        "the embed must be fully internal to the package")

    ctn = _media_node_ctn(slide1)
    assert ctn is not None and ctn.get("fill") == "hold"
    cond = next(el for el in ctn.iter() if _local(el.tag) == "cond")
    assert cond.get("delay") == "indefinite", (
        "on_click keeps the tested indefinite start condition")
    assert ctn.get("repeatCount") is None
    fixture = (HERE / "fixtures" / "mg036_timing_onclick.xml").read_bytes()
    assert _timing_signature(slide1) == _timing_signature(fixture)

    from pptx import Presentation
    prs = Presentation(str(deck))
    assert [s.notes_slide.notes_text_frame.text for s in prs.slides] == [
        SPEECH[1], SPEECH[2], SPEECH[3]], (
        "speaker notes survive the motion embed")

    deck2 = tmp_path / "deck2.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), deck2,
                             speech_chunks=SPEECH,
                             motion_media=motion_media)
    parts2 = _deck_parts(deck2)
    assert len([n for n in parts2 if n.startswith("ppt/media/")
                and n.lower().endswith(".mp4")]) == 1, (
        "re-assembly builds fresh and never duplicates media")


@needs_media
def test_assemble_automatic_loop_timing(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(
        tmp_path, media,
        package_overrides={
            "media": {"fit": "approved_loop", "playback": LOOP_PLAYBACK},
            "package": {"playback": LOOP_PLAYBACK,
                        "loop_seam_status": "pass"}})
    _record, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    assert set(motion_media) == {1}
    deck = tmp_path / "deck-loop.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), deck,
                             motion_media=motion_media)
    slide1 = _deck_parts(deck)["ppt/slides/slide1.xml"]
    ctn = _media_node_ctn(slide1)
    cond = next(el for el in ctn.iter() if _local(el.tag) == "cond")
    assert cond.get("delay") == "0", "automatic start writes delay 0"
    assert ctn.get("repeatCount") == "indefinite", (
        "loop writes repeatCount indefinite on the media node")
    fixture = (HERE / "fixtures"
               / "mg036_timing_automatic_loop.xml").read_bytes()
    assert _timing_signature(slide1) == _timing_signature(fixture)


@needs_media
def test_finalize_marks_embedded_and_missing(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(tmp_path / "a", media)
    _r, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    deck = tmp_path / "a" / "deck.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), deck,
                             motion_media=motion_media)
    record = build_deck.finalize_motion_record(run_dir, deck)
    entry = _outcome_for(record, 1)
    assert entry["status"] == "embedded", entry
    assert entry["embed"]["media_part"].startswith("ppt/media/")

    run_dir2 = build_run(tmp_path / "b", media)
    build_deck.plan_motion_media(run_dir2, _slides_of(run_dir2))
    deck2 = tmp_path / "b" / "deck.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir2), deck2,
                             motion_media=None)
    record2 = build_deck.finalize_motion_record(run_dir2, deck2)
    entry2 = _outcome_for(record2, 1)
    assert entry2["status"] == "degraded"
    assert entry2["reason_code"] == pm.MOTION_EMBED_MISSING


# ---------------------------------------------------------------------------
# P08: the P8.15-MOTION phase audit + its verifier
# ---------------------------------------------------------------------------

def _delivered_deck(run_dir: Path, deck: Path) -> Path:
    target = run_dir / "fixture-FINAL.pptx"
    shutil.copy2(deck, target)
    return target


@needs_media
def test_motion_phase_audit_and_verifier(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(tmp_path, media)
    _r, motion_media = build_deck.plan_motion_media(
        run_dir, _slides_of(run_dir))
    deck = tmp_path / "deck.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), deck,
                             motion_media=motion_media)
    build_deck.finalize_motion_record(run_dir, deck)
    _delivered_deck(run_dir, deck)
    assert motion_phase.run(run_dir) == motion_phase.EXIT_OK
    ok, reasons = phase_verifiers.verify("P8.15-MOTION", run_dir)
    assert ok, reasons
    record = json.loads((run_dir / pm.MOTION_RECORD_REL).read_text())
    assert record["phase_audit"]["embedded_verified"] == 1
    assert record["phase_audit"]["failures"] == []


@needs_media
def test_motion_phase_catches_forged_embed_claim(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(tmp_path, media)
    build_deck.plan_motion_media(run_dir, _slides_of(run_dir))
    still_deck = tmp_path / "still.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), still_deck,
                             motion_media=None)
    # Forge the record: claim the slide embedded though the delivered
    # deck is the still-only assembly.
    record = json.loads((run_dir / pm.MOTION_RECORD_REL).read_text())
    entry = _outcome_for(record, 1)
    entry["status"] = "embedded"
    _write_json(run_dir / pm.MOTION_RECORD_REL, record)
    _delivered_deck(run_dir, still_deck)
    assert motion_phase.run(run_dir) == motion_phase.EXIT_AUDIT_FAILED
    ok, reasons = phase_verifiers.verify("P8.15-MOTION", run_dir)
    assert not ok and reasons


@needs_media
def test_motion_phase_late_piece_is_not_a_failure(media, tmp_path):
    pytest.importorskip("pptx")
    run_dir = build_run(tmp_path, media)
    still_deck = tmp_path / "still.pptx"
    build_deck.assemble_pptx(_rendered_list(run_dir), still_deck,
                             motion_media=None)
    _delivered_deck(run_dir, still_deck)
    # No motion.json exists: assembly ran outside the governed path.
    assert not (run_dir / pm.MOTION_RECORD_REL).exists()
    assert motion_phase.run(run_dir) == motion_phase.EXIT_OK
    record = json.loads((run_dir / pm.MOTION_RECORD_REL).read_text())
    entry = _outcome_for(record, 1)
    assert entry["status"] == "produced_late_not_embedded", entry
    assert record["phase_audit"]["produced_late_not_embedded"] == [1]


def test_motion_phase_missing_inputs_exit_3(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    assert motion_phase.run(empty) == motion_phase.EXIT_INPUT_MISSING


# ---------------------------------------------------------------------------
# P03: the narrow authenticated paths in image QC + canonical render
# ---------------------------------------------------------------------------

def _flat_poster_bytes() -> bytes:
    """A poster that IS a flat light card to the pixel check (over
    90 percent one light fill) yet clears the byte floor through a
    high-entropy strip: exactly the shape only an authenticated
    motion poster may carry past AF-IMAGE-QC-VISION."""
    Image = pytest.importorskip("PIL.Image")
    import io
    import os as _os
    img = Image.new("RGB", (1920, 1080), (247, 244, 238))
    px = img.load()
    for x in range(100, 1300):
        for y in range(200, 260):
            px[x, y] = (192, 74, 38)
    noise = Image.frombytes("RGB", (1920, 54),
                            _os.urandom(1920 * 54 * 3))
    img.paste(noise, (0, 1026))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    data = buf.getvalue()
    assert len(data) >= 51200
    return data


@needs_media
def test_image_qc_vision_exempts_only_authenticated_poster(
        media, tmp_path):
    poster = _flat_poster_bytes()
    run_dir = build_run(tmp_path / "auth", media, poster_bytes=poster)
    assert build_deck.check_image_qc_vision(run_dir) == "", (
        "an authenticated motion poster is exempt from the flat-fill "
        "judgment (never from the byte floor)")

    forged = build_run(tmp_path / "forged", media,
                       poster_bytes=poster)
    shutil.rmtree(forged / "working" / "motion" / "slide-01"
                  / "package")
    result = build_deck.check_image_qc_vision(forged)
    assert "slide-01.png" in result and "single light fill" in result, (
        "the identical PNG without an authenticating package is a "
        "flat card and must be flagged")


@needs_media
def test_canonical_render_path_exempts_only_authenticated_workspace(
        media, tmp_path):
    run_dir = build_run(tmp_path, media)
    hand_rolled = 'from PIL import Image\nimg = Image.new("RGB", (2048, 1152))\n'
    inside = (run_dir / "working" / "motion" / "slide-01"
              / "hand_rolled.py")
    inside.write_text(hand_rolled)
    assert build_deck.check_canonical_render_path(run_dir) == "", (
        "a generator inside an authenticated Skill 72 workspace is "
        "the governed artifact path, not a local canvas")
    outside = (run_dir / "working" / "motion" / "slide-09")
    outside.mkdir(parents=True)
    (outside / "hand_rolled.py").write_text(hand_rolled)
    result = build_deck.check_canonical_render_path(run_dir)
    assert result.startswith("AF-LOCAL-CANVAS"), result
    assert "slide-09" in result


# ---------------------------------------------------------------------------
# P07 + P08 wiring: deliverables, manifest, verifier registry
# ---------------------------------------------------------------------------

def test_deliverables_motion_components_derive_from_canonical_spec(
        monkeypatch):
    assert deliverables.DELIVERABLE_COUNT == 10, (
        "motion adds NO top-level deliverables")
    deps = deliverables.motion_component_dependencies()
    assert set(deps) == {"deck_pptx", "deck_pdf", "webinar_mp4"}
    spec_by_key = {d["key"]: d for d in deliverables.DELIVERABLE_AUDIT_SPEC}
    for key, entry in deps.items():
        assert entry["deliverable"] is spec_by_key[key], (
            "components must hang off the LIVE canonical deliverable, "
            "never a copied spec")
        assert entry["components"], key
        for comp in entry["components"]:
            assert comp["producing_phase"] == "P8.15-MOTION"
            assert all(c.startswith("MOTION_")
                       for c in comp["gate_codes"]), comp
    # If a canonical parent ever leaves the spec, the resolver refuses
    # by name instead of silently orphaning the motion contract.
    monkeypatch.setattr(
        deliverables, "MOTION_COMPONENT_DEPENDENCIES",
        {"deck_pptx": [], "bogus_parent": []})
    with pytest.raises(KeyError, match="bogus_parent"):
        deliverables.motion_component_dependencies()


def test_manifest_and_verifier_carry_the_motion_phase():
    from presentation_job import manifest as pj_manifest
    man = json.loads(PIPELINE_MANIFEST.read_text())
    assert man["manifest_version"] == pj_manifest.MIN_MANIFEST_VERSION
    phase = next(p for p in man["phases"] if p["id"] == "P8.15-MOTION")
    assert phase["executor"]["kind"] == "script"
    assert "motion_phase.py" in phase["executor"]["cmd"]
    assert phase["preflight"]["checker"] == "_chk_motion_media"
    assert hasattr(build_deck, "_chk_motion_media")
    assert "P8.15-MOTION" in pj_manifest.PHASE_BUDGET_MINUTES
    webinar = next(p for p in man["phases"]
                   if p["id"] == "P9.6-WEBINAR-VIDEO")
    assert "working/checkpoints/slide_media_manifest.json" in \
        webinar["consumes"]
    assert phase_verifiers.PHASE_VERIFIERS["P8.15-MOTION"] is \
        phase_verifiers._verify_motion_phase


# ---------------------------------------------------------------------------
# P05: build_webinar_video.resolve_slide_media (hash-verified intake)
# ---------------------------------------------------------------------------

@needs_media
def test_resolve_slide_media_for_webinar(media, tmp_path):
    run_dir = build_run(tmp_path, media)
    build_deck.plan_motion_media(run_dir, _slides_of(run_dir))
    resolved = bwv.resolve_slide_media(run_dir)
    assert set(resolved) == {1}
    assert resolved[1]["fit"] == "hold_final"
    clip = Path(resolved[1]["clip_path"])
    assert clip.is_file() and _sha256_file(clip) == _sha256_file(
        media["clip"])
    # Tamper with the clip: the consumer must refuse by name, never
    # render unverified bytes into the webinar video.
    clip.write_bytes(clip.read_bytes() + b"x")
    with pytest.raises(bwv.WebinarBuildError,
                       match="MOTION_HASH_MISMATCH"):
        bwv.resolve_slide_media(run_dir)
    empty = tmp_path / "empty"
    empty.mkdir()
    assert bwv.resolve_slide_media(empty) == {}


# ---------------------------------------------------------------------------
# P06: the webinar render reuses the SAME component at its slot
# ---------------------------------------------------------------------------

def _tone(path: Path, seconds: float, freq: int = 440) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    _run_ffmpeg(["-f", "lavfi", "-i",
                 f"sine=frequency={freq}:d={seconds}",
                 "-c:a", "libmp3lame", str(path)])
    return path


def _frame_at(video: Path, t: float, out: Path) -> Path:
    _run_ffmpeg(["-ss", f"{t}", "-i", str(video), "-frames:v", "1",
                 "-vf", "scale=320:180", str(out)])
    return out


def _frame_diff(a: Path, b: Path) -> float:
    Image = pytest.importorskip("PIL.Image")
    ia = Image.open(a).convert("L")
    ib = Image.open(b).convert("L")
    assert ia.size == ib.size
    pa, pb = ia.load(), ib.load()
    total = 0
    w, h = ia.size
    for y in range(h):
        for x in range(w):
            total += abs(pa[x, y] - pb[x, y])
    return total / (w * h)


@pytest.fixture(scope="module")
def webinar(tmp_path_factory, media):
    """One timing track rendered twice: with the motion branch and
    still-only. Slide 1's slot is exactly the component's 6.0 s."""
    d = tmp_path_factory.mktemp("mg036-webinar")
    renders = d / "renders"
    renders.mkdir()
    (renders / "slide-01.png").write_bytes(media["poster"].read_bytes())
    _still_png(renders / "slide-02.png", 2)
    _still_png(renders / "slide-03.png", 3)
    timing = _write_json(d / "webinar_timing.json", {"timing": [
        {"slide": 1, "duration": 6.0},
        {"slide": 2, "duration": 5.0},
        {"slide": 3, "duration": 4.0}]})
    audio = _tone(d / "audio.mp3", 15.0)
    common = dict(fps=30.0, fade=0.5, crf=22, preset="veryfast",
                  verbose=False)
    motion_out = d / "webinar-motion.mp4"
    probe_motion = wf.render_webinar(
        str(renders), str(timing), str(audio), str(motion_out),
        slide_media={1: {"clip_path": str(media["clip"]),
                         "fit": "hold_final"}}, **common)
    still_out = d / "webinar-still.mp4"
    probe_still = wf.render_webinar(
        str(renders), str(timing), str(audio), str(still_out), **common)
    return {"dir": d, "renders": renders, "timing": timing,
            "audio": audio, "motion_out": motion_out,
            "still_out": still_out, "probe_motion": probe_motion,
            "probe_still": probe_still}


@needs_media
def test_webinar_motion_render_placement(webinar, media, tmp_path):
    probe = webinar["probe_motion"]
    # 6 + 5 + 4 seconds of slots, two 0.5 s xfades: 14.0 s total, and
    # the parent narration is present (no second audio branch).
    assert probe["duration"] == pytest.approx(14.0, abs=0.4)
    assert probe["has_audio"] and probe["has_video"]
    assert webinar["probe_still"]["duration"] == pytest.approx(
        14.0, abs=0.4), "the motion branch must not move the timeline"

    d = tmp_path / "frames"
    d.mkdir()
    f_motion = _frame_at(webinar["motion_out"], 2.0, d / "m2.png")
    f_comp = _frame_at(media["clip"], 2.0, d / "c2.png")
    f_still = _frame_at(webinar["still_out"], 2.0, d / "s2.png")
    poster = _frame_at(media["clip"], 5.9, d / "poster.png")
    d_motion_comp = _frame_diff(f_motion, f_comp)
    d_motion_poster = _frame_diff(f_motion, poster)
    d_still_poster = _frame_diff(f_still, poster)
    assert d_motion_comp < 25.0, (
        f"the webinar's slide-1 window must BE the component playing "
        f"(mean abs diff vs the component at t=2s: {d_motion_comp:.1f})")
    assert d_motion_poster > 4.0 and \
        d_motion_poster > d_motion_comp + 3.0, (
        "the motion window differs from the poster: the component is "
        "playing, not flattened to its poster")
    assert d_still_poster < d_motion_poster, (
        "the still-only render of the same slot stays at the poster "
        "(Ken Burns), proving the difference is the motion branch")


@needs_media
def test_webinar_exact_slot_refusal(webinar, media):
    with pytest.raises(wf.WebinarError,
                       match="MOTION_CLIP_EXCEEDS_SLOT"):
        _exact_slot_render(webinar, media)


def _exact_slot_render(webinar, media):
    """exact_slot against a 5.0 s slot for the 6.0 s component: the
    render must refuse by name (never truncate, never stretch)."""
    d = webinar["dir"]
    timing = _write_json(d / "timing-5s.json", {"timing": [
        {"slide": 1, "duration": 5.0},
        {"slide": 2, "duration": 5.0},
        {"slide": 3, "duration": 4.0}]})
    return wf.render_webinar(
        str(webinar["renders"]), str(timing), str(webinar["audio"]),
        str(d / "never.mp4"), fps=30.0, fade=0.5, crf=22,
        preset="veryfast",
        slide_media={1: {"clip_path": str(media["clip"]),
                         "fit": "exact_slot"}})


@needs_media
def test_webinar_hold_final_pads_with_final_frame(
        webinar, media, tmp_path):
    d = tmp_path / "hold"
    d.mkdir()
    renders = d / "renders"
    renders.mkdir()
    (renders / "slide-01.png").write_bytes(media["poster"].read_bytes())
    _still_png(renders / "slide-02.png", 2)
    timing = _write_json(d / "timing.json", {"timing": [
        {"slide": 1, "duration": 8.0}, {"slide": 2, "duration": 2.0}]})
    audio = _tone(d / "audio.mp3", 10.0)
    out = d / "out.mp4"
    probe = wf.render_webinar(
        str(renders), str(timing), str(audio), str(out), fps=30.0,
        fade=0.5, crf=22, preset="veryfast",
        slide_media={1: {"clip_path": str(media["clip"]),
                         "fit": "hold_final"}})
    assert probe["duration"] == pytest.approx(9.5, abs=0.4)
    fd = tmp_path / "hold-frames"
    fd.mkdir()
    # 7.4 s: inside the held tail, before the 7.5 s xfade begins.
    f_held = _frame_at(out, 7.4, fd / "held.png")
    poster = _frame_at(media["clip"], 5.9, fd / "poster.png")
    f_mid = _frame_at(media["clip"], 2.0, fd / "mid.png")
    assert _frame_diff(f_held, poster) < 20.0, (
        "past the component's end, hold_final shows its final frame")
    assert _frame_diff(f_held, f_mid) > 8.0


# ---------------------------------------------------------------------------
# Still-only regression: the pre-change pipeline (pinned pre-MG-036
# base, REGRESSION_BASE_REF) vs the edited pipeline on identical
# still inputs
# ---------------------------------------------------------------------------

@needs_media
def test_still_only_regression_vs_base(media, tmp_path):
    pytest.importorskip("pptx")
    git = shutil.which("git")
    if not git:
        pytest.skip("git unavailable for the base comparison")
    rel = ("23-ai-workforce-blueprint/templates/role-library/"
           "presentations/scripts/build_deck.py")
    base_ref = os.environ.get("MOTION_REGRESSION_BASE") \
        or REGRESSION_BASE_REF
    got = subprocess.run([git, "-C", str(REPO), "show",
                          f"{base_ref}:{rel}"],
                         capture_output=True, text=True)
    if got.returncode != 0:
        pytest.skip(
            f"base build_deck.py not retrievable at {base_ref}")
    # The base module resolves shared-utils relative to its own file,
    # so it must be staged inside the scripts tree (and removed right
    # after the import) rather than in a bare tmp dir.
    base_file = SCRIPTS / "_mg036_base_build_deck.py"
    base_file.write_text(got.stdout)
    import importlib.util
    try:
        spec = importlib.util.spec_from_file_location("build_deck_base",
                                                      base_file)
        base_mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(base_mod)
    finally:
        base_file.unlink(missing_ok=True)

    run_dir = build_run(tmp_path / "run", media, motion_slides=())
    rendered = _rendered_list(run_dir)
    base_deck = tmp_path / "base.pptx"
    new_deck = tmp_path / "new.pptx"
    base_mod.assemble_pptx(rendered, base_deck, speech_chunks=SPEECH)
    build_deck.assemble_pptx(rendered, new_deck, speech_chunks=SPEECH,
                             motion_media=None)
    bp, np_ = _deck_parts(base_deck), _deck_parts(new_deck)
    assert sorted(bp) == sorted(np_), (
        "the still path must produce the identical part set")
    assert not any(n.lower().endswith((".mp4", ".mov", ".m4v", ".webm"))
                   for n in np_), "a still deck carries no movie part"
    import xml.etree.ElementTree as ET
    for i in (1, 2, 3):
        for parts, label in ((bp, "base"), (np_, "new")):
            pics = sum(1 for el in ET.fromstring(
                parts[f"ppt/slides/slide{i}.xml"]).iter()
                if _local(el.tag) == "pic")
            assert pics == 1, (label, i, pics)
    from pptx import Presentation
    base_notes = [s.notes_slide.notes_text_frame.text
                  for s in Presentation(str(base_deck)).slides]
    new_notes = [s.notes_slide.notes_text_frame.text
                 for s in Presentation(str(new_deck)).slides]
    assert base_notes == new_notes == [SPEECH[1], SPEECH[2], SPEECH[3]]
